"""Validated readers and content-based discovery for manually replaced RBI surveys.

Only the two canonical bi-monthly folders are consulted. Files are identified
by published table signatures and geography notes, never release filenames.
Unknown/malformed files and tied newest editions are explicit errors.
"""
from __future__ import annotations
from datetime import datetime
import hashlib
import math
from numbers import Real
from pathlib import Path
import re
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
INFLATION_DIR = ROOT / 'data/rbi_bimonthly_manual/inflation_survey'
CONSUMER_DIR = ROOT / 'data/rbi_bimonthly_manual/consumer_survey'
INFLATION_SOURCE = 'https://www.rbi.org.in/Scripts/BimonthlyPublications.aspx?head=Inflation%20Expectations%20Survey%20of%20Households%20-%20Bi-monthly'
CONSUMER_SOURCES = {'Urban': 'https://www.rbi.org.in/Scripts/BimonthlyPublications.aspx?head=Urban%20Consumer%20Confidence%20Survey%20-%20Bi-monthly', 'Rural': 'https://www.rbi.org.in/Scripts/BimonthlyPublications.aspx?head=Rural%20Consumer%20Confidence%20Survey%20-%20Bi-monthly'}

SHEET = "T2 Household IE"
SERIES = {"current": 5, "three_month": 8, "one_year": 11}

SHEETS = {"three_month": "T1A Product-wise 3M", "one_year": "T1B Product-wise 1Y"}
CATEGORIES = {"Food": "Options: Food Product", "Non-food": "Options: Non- Food Product",
              "Housing": "Options: Housing Prices", "Services": "Options: Cost of Services",
              "Household durables": "Options: Household Durables"}

COMPONENTS=[('Economic conditions','general economic situation'),
            ('Employment','expectations on employment'),('Prices','expectations on price level'),
            ('Income','expectations on income'),('Spending','expectations on spending')]

def read_observations(path=None):
    """Select dated observations only; bracketed standard errors are separate rows."""
    path = discover_inflation() if path is None else Path(path)
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet = workbook[SHEET]
        expected = {"B4": "Round No.", "C4": "Survey period ended",
                    "D5": "Current", "G5": "3 months ahead", "J5": "1 year ahead",
                    "E6": "Median", "H6": "Median", "K6": "Median"}
        for cell, label in expected.items():
            if str(sheet[cell].value).strip() != label:
                raise ValueError(f"Unexpected header at {cell}; expected {label}")
        records, missing = [], []
        for row in range(7, sheet.max_row + 1):
            date = sheet.cell(row, 3).value
            round_no = sheet.cell(row, 2).value
            if not isinstance(date, datetime):
                # Metadata and merged standard-error rows have no round/date.
                if isinstance(round_no, Real) or (isinstance(round_no, str)
                        and round_no.strip().rstrip("B").isdigit()):
                    raise ValueError(f"Missing/invalid survey date at C{row}")
                continue
            if round_no is None:
                raise ValueError(f"Missing round at B{row}")
            record = {"date": date, "round": str(round_no), "source_row": row}
            for name, column in SERIES.items():
                value = sheet.cell(row, column).value
                if value is None:
                    missing.append({"row": row, "series": name})
                    record[name] = math.nan
                elif isinstance(value, bool) or not isinstance(value, Real) or not math.isfinite(value):
                    raise ValueError(f"Invalid numeric observation at row {row}, column {column}")
                else:
                    record[name] = float(value)
            records.append(record)
        if len(records) < 3:
            raise ValueError("At least three survey observations are required")
        if any(a["date"] >= b["date"] for a, b in zip(records, records[1:])):
            raise ValueError("Survey dates must be unique and strictly chronological")
        if len({r["round"] for r in records}) != len(records):
            raise ValueError("Duplicate survey rounds")
        return records, missing
    finally:
        workbook.close()

def select_curves(records):
    latest, previous = records[-1], records[-2]
    target = latest["date"].replace(year=latest["date"].year - 1, day=min(latest["date"].day, 28))
    earlier = min(records[:-1], key=lambda r: abs((r["date"] - target).days))
    if abs((earlier["date"] - target).days) > 100:
        raise ValueError("No survey within 100 days of the year-earlier target")
    for record in (latest, previous, earlier):
        if not all(math.isfinite(record[k]) for k in SERIES):
            raise ValueError(f"Missing comparison data for {record['date']:%b %Y}; no substitution")
    return latest, previous, earlier

def serializable(record):
    return {key: (value.isoformat()[:10] if isinstance(value, datetime)
                  else None if isinstance(value, float) and not math.isfinite(value)
                  else value) for key, value in record.items()}

def read_categories(path=None):
    """Use the latest dated estimate column in each table, never SE or coherence."""
    path = discover_inflation() if path is None else Path(path)
    workbook = openpyxl.load_workbook(path, data_only=True)
    records = {name: {"category": name, "source_category": label} for name, label in CATEGORIES.items()}
    periods = []
    try:
        for horizon, name in SHEETS.items():
            sheet = workbook[name]
            title = str(sheet["B1" if horizon == "three_month" else "B2"].value)
            expected_horizon = "Three months ahead" if horizon == "three_month" else "One year ahead"
            if expected_horizon not in title or sheet["B3"].value != "(Percentage of Respondents)":
                raise ValueError(f"Unexpected table definition: {name}")
            if sheet["B4"].value != "Round No." or sheet["B5"].value != "Survey period ended":
                raise ValueError(f"Unexpected round/date headers: {name}")
            dated = [cell for cell in sheet[5] if isinstance(cell.value, datetime)]
            if not dated:
                raise ValueError(f"No dated survey columns: {name}")
            date = max(cell.value for cell in dated)
            latest = [cell for cell in dated if cell.value == date]
            if len(latest) != 1:
                raise ValueError(f"Duplicate latest survey: {name}")
            column = latest[0].column
            round_no = sheet.cell(4, column).value
            if round_no is None or sheet.cell(6, column).value != "Estimate":
                raise ValueError(f"Missing round or not an estimate column: {name}")
            periods.append((date, str(round_no)))
            for category, label in CATEGORIES.items():
                rows = [r for r in range(1, sheet.max_row + 1) if sheet.cell(r, 2).value == label]
                if len(rows) != 1 or sheet.cell(rows[0] + 1, 2).value != "Prices will increase":
                    raise ValueError(f"Missing/ambiguous category or changed measure: {category}")
                row = rows[0] + 1
                cell = sheet.cell(row, column)
                value = cell.value
                if isinstance(value, bool) or not isinstance(value, Real) or not math.isfinite(value) or not 0 <= value <= 100:
                    raise ValueError(f"Invalid percentage: {name}!{cell.coordinate}")
                # Independently reconcile with the three price-increase subresponses.
                parts = [sheet.cell(row + n, column).value for n in (1, 2, 3)]
                if any(not isinstance(v, Real) or isinstance(v, bool) or not math.isfinite(v) for v in parts):
                    raise ValueError(f"Missing price-increase components: {category}")
                if abs(sum(parts) - value) > .05:
                    raise ValueError(f"Price-increase components do not reconcile: {category}")
                records[category][horizon] = float(value)
                records[category][horizon + "_cell"] = f"{name}!{cell.coordinate}"
        if periods[0] != periods[1]:
            raise ValueError("Latest horizon rounds differ; no substitution")
        for record in records.values():
            record["change_pp"] = record["one_year"] - record["three_month"]
        ordered = sorted(records.values(), key=lambda r: -r["change_pp"])
        return {"survey_date": periods[0][0].strftime("%Y-%m-%d"), "round": periods[0][1],
                "sheets": SHEETS, "measure": "Prices will increase", "units": "Percentage of respondents",
                "ordering": "Descending unrounded 1Y minus 3M percentage-point change",
                "categories": ordered, "workbook": str(Path(path).resolve()),
                "workbook_sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest()}
    finally:
        workbook.close()

def norm(value):
    return re.sub(r'\s+', ' ', str(value or '').strip()).lower()

def number(value, location, low, high):
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not low <= value <= high:
        raise ValueError(f'Invalid value at {location}: {value!r}')
    return float(value)

def find_table(workbook, phrase):
    matches = [s for s in workbook if s.sheet_state == 'visible' and
               any(phrase in norm(c.value) for row in s for c in row)]
    if len(matches) != 1:
        raise ValueError(f'Expected one published table for {phrase!r}: {[s.title for s in matches]}')
    return matches[0]

def _read_table(sheet, indices=False):
    """Discover Survey Round and statistic columns, including merged horizon headers."""
    headers = [c for row in sheet for c in row if norm(c.value) == 'survey round']
    if len(headers) != 1:
        raise ValueError(f'Ambiguous date header in {sheet.title}')
    header = headers[0]
    if indices:
        columns = {k: next(c.column for c in sheet[header.row] if norm(c.value).rstrip('*') == k)
                   for k in ('csi', 'fei')}
    else:
        horizons = {'current': 'current perception', 'ahead': 'one year ahead expectation'}
        columns = {}
        for key, label in horizons.items():
            group = next(c for c in sheet[header.row] if norm(c.value) == label)
            merged = next((r for r in sheet.merged_cells.ranges if group.coordinate in r), None)
            if merged is None:
                raise ValueError(f'Missing horizon structure: {sheet.title}')
            sub = {norm(sheet.cell(header.row + 1, col).value): col
                   for col in range(merged.min_col, merged.max_col + 1)}
            for statistic, labels in {'net': ['net response'], 'increase': ['increased', 'will increase'],
                                      'same': ['remained same', 'will remain same'],
                                      'decrease': ['decreased', 'will decrease']}.items():
                columns[f'{key}_{statistic}'] = next(sub[x] for x in labels if x in sub)
    records = []
    for row in sheet.iter_rows(min_row=header.row + 1):
        date = row[header.column - 1].value
        if not isinstance(date, datetime):
            if any(isinstance(row[col-1].value, (int, float)) for col in columns.values()):
                raise ValueError(f'Missing or invalid survey date: {sheet.title}!{row[header.column-1].coordinate}')
            continue
        record = {'date': date.strftime('%Y-%m-%d'), 'source_date_cell': row[header.column - 1].coordinate}
        for key, col in columns.items():
            cell = row[col - 1]
            record[key] = number(cell.value, f'{sheet.title}!{cell.coordinate}',
                                 0 if indices or not key.endswith('_net') else -100,
                                 200 if indices else 100)
            record[f'{key}_cell'] = cell.coordinate
        records.append(record)
    dates = [r['date'] for r in records]
    if not dates or dates != sorted(set(dates)):
        raise ValueError(f'Duplicate or unordered survey dates: {sheet.title}')
    return records, {k: openpyxl.utils.get_column_letter(c) for k, c in columns.items()}

def read_table(sheet, indices=False):
    try:
        return _read_table(sheet, indices)
    except (StopIteration, KeyError) as exc:
        raise ValueError(f'Missing required survey date, horizon or statistic header in {sheet.title}') from exc


def load_consumer(path):
    path = Path(path)
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        indices = find_table(workbook, 'current situation index (csi) and future expectations index (fei)')
        essential = find_table(workbook, 'spending- essential items')
        discretionary = find_table(workbook, 'spending- non-essential items')
        notes = [{'sheet': s.title, 'cell': c.coordinate, 'text': c.value}
                 for s in (indices, essential, discretionary) for row in s for c in row
                 if isinstance(c.value, str) and any(x in norm(c.value) for x in
                     ('100 + average', 'coverage', 'states/ uts', 'web releases', 'rural and semi-urban'))]
        if not any('100 + average of net responses' in norm(n['text']) for n in notes):
            raise ValueError('Workbook does not confirm index definition')
        texts = [norm(cell.value) for sheet in workbook if sheet.sheet_state == 'visible'
                 for row in sheet for cell in row if isinstance(cell.value, str)]
        urban = any(re.search(r'\buccs\b|\burban consumer confidence', t) for t in texts)
        rural = any('rural consumer confidence' in t or 'rural and semi-urban' in t for t in texts)
        if urban == rural:
            raise ValueError(f'Cannot uniquely identify Urban/Rural survey from workbook contents: {path.name}')
        geography = 'Urban' if urban else 'Rural'
        result = {'geography': geography, 'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                  'notes': notes, 'inventory': [{'sheet': s.title, 'state': s.sheet_state,
                      'rows': s.max_row, 'columns': s.max_column} for s in workbook], 'tables': {}}
        for key, sheet in [('indices', indices), ('essential', essential), ('nonessential', discretionary)]:
            rows, cols = read_table(sheet, key == 'indices')
            result['tables'][key] = {'sheet': sheet.title, 'columns': cols, 'records': rows}
        return result
    finally:
        workbook.close()

def compare(surveys):
    urban, rural = (surveys[g]['tables']['indices']['records'] for g in ('Urban', 'Rural'))
    common = sorted({r['date'] for r in urban} & {r['date'] for r in rural})
    if len(common) < 2:
        raise ValueError('Insufficient common history')
    latest = common[-1]
    # Require simultaneous latest rounds; never silently show an older snapshot.
    if any(s['tables'][k]['records'][-1]['date'] != latest for s in surveys.values()
           for k in ('indices', 'essential', 'nonessential')):
        raise ValueError('Latest survey rounds are not aligned')
    series, spending = {}, []
    for g, survey in surveys.items():
        lookup = {r['date']: r for r in survey['tables']['indices']['records']}
        series[g] = [lookup[d] for d in common]
        if any(series[g][-1][k] is None for k in ('csi', 'fei')):
            raise ValueError('Missing latest index; no substitution')
    reconciliation = []
    for horizon in ('current', 'ahead'):
        for category in ('essential', 'nonessential'):
            row = {'horizon': horizon, 'category': category}
            for g, survey in surveys.items():
                r = survey['tables'][category]['records'][-1]
                values = [r[f'{horizon}_{k}'] for k in ('net', 'increase', 'same', 'decrease')]
                if any(v is None for v in values):
                    raise ValueError('Missing latest spending; no substitution')
                net, inc, same, dec = values
                if abs(net - (inc - dec)) > .151 or abs(inc + same + dec - 100) > .151:
                    raise ValueError('Latest shares and RBI net response do not reconcile within rounding')
                row[g] = net  # RBI-provided statistic; inc-dec is a validation ONLY.
                reconciliation.append({'geography': g, 'table': category, 'horizon': horizon,
                    'net_cell': r[f'{horizon}_net_cell'], 'reported_net': net,
                    'increase': inc, 'same': same, 'decrease': dec,
                    'rounding_residual': net - (inc - dec)})
            spending.append(row)
    return common, series, spending, reconciliation

def coverage_break(survey):
    texts = [norm(n['text']) for n in survey['notes']]
    confirmed = any('26 states/ uts' in t and '31 states/ uts' in t and 'from july 2024' in t for t in texts)
    return datetime(2024, 7, 1) if confirmed else None

def _read_components(paths=None):
    """Separate reader for component tables; the existing spending/index parser is unchanged.

    RBI already defines prices as decrease minus increase (favourable minus
    unfavourable). Read its Net Response directly; NEVER reverse its sign.
    Overall spending follows RBI's index convention (increase minus decrease),
    which is an expenditure assessment, not a statement about real volumes.
    """
    paths = discover_consumer() if paths is None else paths
    audit=[]; datesets=[]; sources={}
    for g in ('Urban', 'Rural'):
        path=Path(paths[g]); w=openpyxl.load_workbook(path,data_only=True)
        try:
            sources[g]={'url':CONSUMER_SOURCES[g], 'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
            for name,phrase in COMPONENTS:
                matches=[s for s in w if s.sheet_state=='visible' and any(
                    norm(cell.value).startswith('table ') and norm(cell.value).endswith(phrase)
                    for row in s for cell in row)]
                if len(matches)!=1:raise ValueError(f'Ambiguous published component table: {phrase}')
                s=matches[0]
                title=next(cell.value for row in s for cell in row
                           if isinstance(cell.value,str) and phrase in norm(cell.value))
                header=next(cell for row in s for cell in row if norm(cell.value)=='survey round')
                cols={}
                for horizon,label in [('current','current perception'),('ahead','one year ahead expectation')]:
                    group=next(cell for cell in s[header.row] if norm(cell.value)==label)
                    merged=next(r for r in s.merged_cells.ranges if group.coordinate in r)
                    sub={norm(s.cell(header.row+1,col).value):col for col in range(merged.min_col,merged.max_col+1)}
                    up=next(sub[t] for t in ('improved','will improve','increased','will increase') if t in sub)
                    down=next(sub[t] for t in ('worsened','will worsen','decreased','will decrease') if t in sub)
                    same=next(sub[t] for t in ('remained same','will remain same') if t in sub)
                    cols[horizon]=(sub['net response'],up,down,same)
                tabledates=[]
                for row in s.iter_rows(min_row=header.row+2):
                    date=row[header.column-1].value
                    if not isinstance(date,datetime):
                        if any(isinstance(row[col-1].value,(int,float)) for cs in cols.values() for col in cs):
                            raise ValueError(f'Invalid component survey date in {s.title}')
                        continue
                    d=date.strftime('%Y-%m-%d');tabledates.append(d)
                    for horizon,(net,up,down,same) in cols.items():
                        cells=[row[col-1] for col in (net,up,down,same)]
                        vals=[number(cell.value,f'{s.title}!{cell.coordinate}',-100 if i==0 else 0,100)
                              for i,cell in enumerate(cells)]
                        audit.append({'geography':g,'date':d,'horizon':horizon,'component':name,
                            'workbook':path.name,'sheet':s.title,'table':title,
                            'raw_statistic':'Net Response','raw_cell':cells[0].coordinate,
                            'raw_balance':vals[0],'supplied_or_calculated':'RBI supplied',
                            'increase_or_improve':vals[1],'decrease_or_worsen':vals[2],
                            'same':vals[3], 'sign_reversal':False,'plotted_value':vals[0],
                            'orientation':'decrease minus increase' if name=='Prices' else 'increase/improve minus decrease/worsen',
                            'units':'percentage points'})
                if not tabledates or tabledates!=sorted(set(tabledates)):
                    raise ValueError(f'Unordered/duplicate component dates in {s.title}')
                datesets.append(set(tabledates))
        finally: w.close()
    common=sorted(set.intersection(*datesets))
    if len(common)<2: raise ValueError('Two common component rounds are required')
    latest=common[-1]
    previous=common[-2]
    if any(max(ds)!=latest for ds in datesets): raise ValueError('Latest component rounds differ')
    pairedrows=[r for r in audit if r['date'] in (latest,previous)]
    for r in pairedrows:
        if any(r[k] is None for k in ('raw_balance','increase_or_improve','decrease_or_worsen','same')):
            raise ValueError('Missing latest/previous component; no substitution')
        expected=r['decrease_or_worsen']-r['increase_or_improve'] if r['component']=='Prices' else r['increase_or_improve']-r['decrease_or_worsen']
        r['rounding_residual']=r['raw_balance']-expected
        if abs(r['rounding_residual'])>.151 or abs(r['increase_or_improve']+r['decrease_or_worsen']+r['same']-100)>.151:
            raise ValueError(f'Component orientation/share reconciliation failed: {r}')
    latestrows=[r for r in pairedrows if r['date']==latest]
    prev={(r['geography'],r['horizon'],r['component']):r for r in pairedrows if r['date']==previous}
    for r in latestrows:
        before=prev[r['geography'],r['horizon'],r['component']]
        r.update(previous_date=previous,previous_value=before['plotted_value'],
                 previous_raw_cell=before['raw_cell'],previous_raw_statistic=before['raw_statistic'],
                 previous_sign_reversal=before['sign_reversal'],previous_orientation=before['orientation'],
                 previous_rounding_residual=before['rounding_residual'],
                 previous_increase_or_improve=before['increase_or_improve'],
                 previous_decrease_or_worsen=before['decrease_or_worsen'],previous_same=before['same'],
                 delta=r['plotted_value']-before['plotted_value'])
    return latest,previous,latestrows,sources

def read_components(paths=None):
    try:
        return _read_components(paths)
    except (StopIteration, KeyError) as exc:
        raise ValueError('Missing required consumer component table, horizon or Net Response header') from exc


def workbook_files(directory):
    directory = Path(directory)
    paths = sorted(p for p in directory.glob('*') if p.is_file()
                   and p.suffix.lower() == '.xlsx' and not p.name.startswith('~$'))
    if not paths:
        raise ValueError(f'No RBI survey .xlsx workbooks in {directory}')
    return paths


def newest(candidates, geography):
    latest = max(date for date, path in candidates)
    winners = [path for date, path in candidates if date == latest]
    if len(winners) != 1:
        raise ValueError(f'Ambiguous newest {geography} survey ({latest:%Y-%m-%d}): '
                         + ', '.join(p.name for p in winners))
    return winners[0]


def discover_inflation(directory=INFLATION_DIR):
    candidates = []
    for path in workbook_files(directory):
        try:
            records, _ = read_observations(path)
            latest, _, _ = select_curves(records)
            categories = read_categories(path)
            if categories['survey_date'] != latest['date'].strftime('%Y-%m-%d') or categories['round'] != latest['round']:
                raise ValueError('Median and product tables have different latest rounds')
            candidates.append((latest['date'], path))
        except Exception as exc:
            raise ValueError(f'Invalid inflation survey workbook {path.name}: {exc}') from exc
    return newest(candidates, 'inflation')


def discover_consumer(directory=CONSUMER_DIR):
    candidates = {'Urban': [], 'Rural': []}
    for path in workbook_files(directory):
        try:
            survey = load_consumer(path)
            latest = survey['tables']['indices']['records'][-1]['date']
            if any(table['records'][-1]['date'] != latest for table in survey['tables'].values()):
                raise ValueError('Consumer table latest rounds differ')
            candidates[survey['geography']].append((datetime.fromisoformat(latest), path))
        except Exception as exc:
            raise ValueError(f'Invalid consumer survey workbook {path.name}: {exc}') from exc
    for geography, editions in candidates.items():
        if not editions:
            raise ValueError(f'Missing {geography} consumer-confidence workbook in {directory}')
    return {g: newest(editions, g) for g, editions in candidates.items()}


def load_consumer_comparison(directory=CONSUMER_DIR):
    paths = discover_consumer(directory)
    surveys = {g: load_consumer(path) for g, path in paths.items()}
    common, series, spending, reconciliation = compare(surveys)
    latest, previous, rows, sources = read_components(paths)
    if latest != common[-1] or previous != common[-2]:
        raise ValueError('Component and index latest/previous common rounds differ')
    index_reconciliation = []
    for g in ('Urban', 'Rural'):
        for h, k in [('current', 'csi'), ('ahead', 'fei')]:
            reconstructed = 100 + sum(r['plotted_value'] for r in rows
                                     if r['geography'] == g and r['horizon'] == h) / 5
            index = series[g][-1][k]
            if abs(reconstructed-index) > .151:
                raise ValueError(f'{g} {k}: components do not reconcile with published index')
            index_reconciliation.append(dict(geography=g, index=k, reported=index,
                reconstructed=reconstructed, rounding_residual=reconstructed-index))
    return dict(surveys=surveys, common_dates=common, series=series, spending=spending,
                reconciliation=reconciliation, latest_date=latest, previous_date=previous,
                components=rows, sources=sources, index_reconciliation=index_reconciliation,
                coverage_marker=coverage_break(surveys['Rural']))

def load_discretionary_comparison(directory=CONSUMER_DIR, surveys=None):
    """Published Table 8 balances, reconciled for every common spending round."""
    if surveys is None:
        surveys = {g: load_consumer(p) for g, p in discover_consumer(directory).items()}
    lookups = {g: {r['date']: r for r in s['tables']['nonessential']['records']}
               for g, s in surveys.items()}
    common = sorted(set(lookups['Urban']) & set(lookups['Rural']))
    if len(common) < 2:
        raise ValueError('Insufficient common discretionary spending history')
    for g, records in lookups.items():
        if max(records) != common[-1]:
            raise ValueError('Latest discretionary spending rounds are not aligned')
        if {d for d in records if common[0] <= d <= common[-1]} != set(common):
            raise ValueError('Missing common discretionary spending round; no interpolation')
    series = {}
    for g, records in lookups.items():
        series[g] = [records[d] for d in common]
        for row in series[g]:
            for horizon in ('current', 'ahead'):
                values = [row[horizon + '_' + k] for k in ('net', 'increase', 'same', 'decrease')]
                if any(v is None or not math.isfinite(v) for v in values):
                    raise ValueError('Missing discretionary spending observation; no interpolation')
                net, inc, same, dec = values
                if abs(net - (inc - dec)) > .151 or abs(inc + same + dec - 100) > .151:
                    raise ValueError('Discretionary shares and published net response do not reconcile')
    return dict(common_dates=common, latest_date=common[-1], previous_date=common[-2],
                statistic='RBI supplied Net Response: increased minus decreased (pp)',
                series=series, source_urls=CONSUMER_SOURCES,
                sources={g: dict(path=s['path'], sha256=s['sha256'], url=CONSUMER_SOURCES[g],
                                 sheet=s['tables']['nonessential']['sheet'],
                                 columns=s['tables']['nonessential']['columns'])
                         for g, s in surveys.items()})

