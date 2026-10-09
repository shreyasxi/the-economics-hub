"""Manual replacement contract: synthetic releases never enter production output."""
from datetime import datetime
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
from data.processors import rbi_surveys as rbi
from charts.india_charts import india_inflation_surveys as inflation
from charts.india_charts import india_consumer_surveys as consumer
from charts import loader
from config.insights import get_insight


def inflation_fixture(path, latest=datetime(2026, 7, 1), value=8.5):
    w = openpyxl.Workbook(); s = w.active; s.title = rbi.SHEET
    for cell, text in {'B4':'Round No.', 'C4':'Survey period ended',
                       'D5':'Current', 'G5':'3 months ahead', 'J5':'1 year ahead',
                       'E6':'Median', 'H6':'Median', 'K6':'Median'}.items(): s[cell]=text
    dates = [datetime(2025,7,1), datetime(2026,5,1), datetime(2026,7,1)]
    if latest > dates[-1]: dates.append(latest)
    for row,date in enumerate(dates,7):
        s.cell(row,2,str(row)); s.cell(row,3,date)
        for col,offset in [(5,0),(8,1),(11,.5)]: s.cell(row,col,value+offset if date==latest else 7+offset)
    for horizon,name in rbi.SHEETS.items():
        s=w.create_sheet(name)
        s['B1' if horizon=='three_month' else 'B2']='Three months ahead' if horizon=='three_month' else 'One year ahead'
        for cell,text in {'B3':'(Percentage of Respondents)','B4':'Round No.','B5':'Survey period ended'}.items():s[cell]=text
        s['C4']=str(6+len(dates));s['C5']=latest;s['C6']='Estimate'
        for row,label in enumerate(rbi.CATEGORIES.values()):
            base=8+row*6;s.cell(base,2,label);s.cell(base+1,2,'Prices will increase')
            total=70+row+(2 if horizon=='one_year' else 0)
            for offset,val in enumerate([total,total-20,10,10],1):s.cell(base+offset,3,val)
    w.save(path);w.close()


def consumer_fixture(path, geography, latest=datetime(2026,7,1), shift=0):
    w=openpyxl.Workbook();w.remove(w.active)
    dates=[datetime(2025,9,1),datetime(2026,5,1),datetime(2026,7,1)]
    if latest>dates[-1]:dates.append(latest)
    current=[-10,-20,-80,10,60];ahead=[10,20,-60,40,70]
    for component,phrase in rbi.COMPONENTS+[('Essential','spending- essential items'),('Nonessential','spending- non-essential items')]:
        s=w.create_sheet(component);s['B2']='Table 1: Perceptions and '+phrase
        s['B4']='Survey Round';s['C4']='Current Perception';s['G4']='One year ahead Expectation'
        s.merge_cells('C4:F4');s.merge_cells('G4:J4')
        up,down=('Improved','Worsened') if component in ('Economic conditions','Employment') else ('Increased','Decreased')
        for col,text in [(3,up),(4,'Remained Same'),(5,down),(6,'Net Response'),
                         (7,'Will '+up.lower().replace('improved','improve').replace('increased','increase')),
                         (8,'Will Remain Same'),(9,'Will '+down.lower().replace('worsened','worsen').replace('decreased','decrease')),
                         (10,'Net Response')]:s.cell(5,col,text)
        i=next((i for i,(n,_) in enumerate(rbi.COMPONENTS) if n==component),None)
        for row,date in enumerate(dates,6):
            s.cell(row,2,date)
            for first,values in [(3,current),(7,ahead)]:
                net=(values[i] if i is not None else 40)+(shift if date==latest else 0)
                oriented=-net if component=='Prices' else net
                inc=(100+oriented)/2;dec=(100-oriented)/2
                for col,val in zip(range(first,first+4),[inc,0,dec,net]):s.cell(row,col,val)
        s.cell(6+len(dates),2,'Note: UCCS centres' if geography=='Urban' else 'The rural consumer confidence survey covers rural and semi-urban areas')
    s=w.create_sheet('Indices');s['B2']='Table 9: Current Situation Index (CSI) and Future Expectations Index (FEI)'
    s['B3']='Survey Round';s['C3']='CSI*';s['D3']='FEI*'
    for row,date in enumerate(dates,4):
        s.cell(row,2,date);s.cell(row,3,100+sum(current)/5+(shift if date==latest else 0));s.cell(row,4,100+sum(ahead)/5+(shift if date==latest else 0))
    s.cell(4+len(dates),2,'100 + average of net responses')
    w.save(path);w.close()


class ReplacementTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.inf=self.root/'data/inputs/rbi/rbi_bimonthly_manual/inflation_survey';self.inf.mkdir(parents=True)
        self.con=self.root/'data/inputs/rbi/rbi_bimonthly_manual/consumer_survey';self.con.mkdir(parents=True)
        self.output=self.root/'output/india/2026-10'
        inflation_fixture(self.inf/'release.xlsx')
        # Deliberately misleading filenames: content determines geography.
        consumer_fixture(self.con/'RCCSH-arbitrary.xlsx','Urban')
        consumer_fixture(self.con/'UCCSH-arbitrary.xlsx','Rural')

    def test_same_filename_replacement_advances_all_six_charts(self):
        self.assertEqual(len(inflation.generate(self.output,self.inf)),2)
        self.assertEqual(len(consumer.generate(self.output,self.con)),4)
        before={name:json.loads((self.output/(name+'.json')).read_text()) for name in inflation.NAMES+consumer.NAMES}
        next_date=datetime(2026,9,1)
        inflation_fixture(self.inf/'release.xlsx',next_date,9.5)
        consumer_fixture(self.con/'RCCSH-arbitrary.xlsx','Urban',next_date,1)
        consumer_fixture(self.con/'UCCSH-arbitrary.xlsx','Rural',next_date,2)
        self.assertEqual(len(inflation.generate(self.output,self.inf)),2)
        self.assertEqual(len(consumer.generate(self.output,self.con)),4)
        after={name:json.loads((self.output/(name+'.json')).read_text()) for name in inflation.NAMES+consumer.NAMES}
        median=after[inflation.NAMES[0]]
        self.assertEqual(median['latest']['date'],'2026-09-01');self.assertEqual(median['previous']['date'],'2026-07-01')
        self.assertEqual(median['latest']['current'],9.5)
        self.assertNotEqual(median['latest']['current'],before[inflation.NAMES[0]]['latest']['current'])
        self.assertEqual(after[inflation.NAMES[1]]['survey_date'],'2026-09-01')
        for name in consumer.NAMES:
            self.assertEqual(after[name]['latest_date'],'2026-09-01');self.assertEqual(after[name]['previous_date'],'2026-07-01')
        for geography, value in [('Urban', 41), ('Rural', 42)]:
            row = after[consumer.NAMES[3]]['series'][geography][-1]
            self.assertEqual(row['date'], '2026-09-01')
            self.assertEqual(row['current_net'], value)
            self.assertEqual(row['ahead_net'], value)
        self.assertEqual(after[consumer.NAMES[0]]['latest_indices']['Urban']['csi'],93)
        self.assertEqual(after[consumer.NAMES[0]]['latest_indices']['Rural']['csi'],94)
        for name, delta in zip(consumer.NAMES[1:], (1, 2)):
            for row in after[name]['components']:
                self.assertEqual(row['delta'], delta)
        with patch.object(loader,'_get_base_dirs',return_value=[self.root/'assets',self.root/'output']):
            charts,_=loader.get_charts('india')
        self.assertEqual({p.stem for p in charts},set(inflation.NAMES+consumer.NAMES))

    def test_retired_matrix_removed_and_never_loaded(self):
        self.output.mkdir(parents=True)
        for ext in ('.png', '.json'):
            (self.output / (consumer.RETIRED_NAME + ext)).write_text('stale')
        self.assertEqual(len(consumer.generate(self.output, self.con)), 4)
        self.assertFalse(any(self.output.glob(consumer.RETIRED_NAME + '.*')))
        # Archived folders can still contain a retired artifact.
        (self.output / (consumer.RETIRED_NAME + '.png')).write_text('stale')
        with patch.object(loader, '_get_base_dirs', return_value=[self.root / 'output']):
            charts, _ = loader.get_charts('india')
        self.assertNotIn(consumer.RETIRED_NAME, [p.stem for p in charts])

    def test_consumer_insights_use_geography_sources(self):
        combined = get_insight(consumer.NAMES[0] + '.png')
        for geography, name in zip(('Urban', 'Rural'), consumer.NAMES[1:]):
            url = rbi.CONSUMER_SOURCES[geography]
            self.assertIn(url, combined)
            self.assertIn(url, get_insight(name + '.png'))
        self.assertIsNone(get_insight(consumer.RETIRED_NAME + '.png'))

    def test_newest_content_date_wins_over_name_and_mtime(self):
        inflation_fixture(self.inf/'aaa.xlsx',datetime(2026,9,1))
        self.assertEqual(rbi.discover_inflation(self.inf).name,'aaa.xlsx')
        consumer_fixture(self.con/'earlier-alphabetically.xlsx','Urban',datetime(2026,9,1))
        self.assertEqual(rbi.discover_consumer(self.con)['Urban'].name,'earlier-alphabetically.xlsx')

    def test_geography_identified_from_content(self):
        found=rbi.discover_consumer(self.con)
        self.assertTrue(found['Urban'].name.startswith('RCCSH'))
        self.assertTrue(found['Rural'].name.startswith('UCCSH'))

    def test_missing_ambiguous_invalid_and_misaligned_fail(self):
        with self.assertRaisesRegex(ValueError,'No RBI survey'):rbi.discover_inflation(self.root/'missing')
        shutil.copy2(self.inf/'release.xlsx',self.inf/'duplicate.xlsx')
        with self.assertRaisesRegex(ValueError,'Ambiguous newest'):rbi.discover_inflation(self.inf)
        shutil.copy2(self.con/'RCCSH-arbitrary.xlsx',self.con/'duplicate.xlsx')
        with self.assertRaisesRegex(ValueError,'Ambiguous newest Urban'):rbi.discover_consumer(self.con)
        (self.con/'duplicate.xlsx').unlink()
        consumer_fixture(self.con/'RCCSH-arbitrary.xlsx','Urban',datetime(2026,9,1))
        with self.assertRaisesRegex(ValueError,'not aligned'):rbi.load_consumer_comparison(self.con)
        (self.con/'UCCSH-arbitrary.xlsx').unlink()
        with self.assertRaisesRegex(ValueError,'Missing Rural'):rbi.discover_consumer(self.con)
        w=openpyxl.Workbook();w.save(self.con/'unknown.xlsx');w.close()
        with self.assertRaisesRegex(ValueError,'Invalid consumer'):rbi.discover_consumer(self.con)

    def test_bad_price_sign_and_changed_schema_are_rejected(self):
        p=self.con/'RCCSH-arbitrary.xlsx';w=openpyxl.load_workbook(p)
        w['Prices']['F8']=80;w.save(p);w.close()
        with self.assertRaisesRegex(ValueError,'orientation/share'):rbi.load_consumer_comparison(self.con)
        p=self.inf/'release.xlsx';w=openpyxl.load_workbook(p)
        w[rbi.SHEET]['E6']='Mean';w.save(p);w.close()
        with self.assertRaisesRegex(ValueError,'header'):rbi.discover_inflation(self.inf)

    def test_ambiguous_geography_rejected(self):
        p=self.con/'RCCSH-arbitrary.xlsx';w=openpyxl.load_workbook(p)
        w['Indices']['B20']='rural consumer confidence';w.save(p);w.close()
        with self.assertRaisesRegex(ValueError,'uniquely identify'):rbi.discover_consumer(self.con)

    def test_no_old_source_fallback_and_failed_run_removes_stale_charts(self):
        old=self.root/'data'/('rbi_'+'survey');old.mkdir()
        shutil.copy2(self.inf/'release.xlsx',old/'valid.xlsx')
        (self.inf/'release.xlsx').unlink()
        self.output.mkdir(parents=True)
        for name in inflation.NAMES+consumer.NAMES:
            for ext in ('.png','.json'):(self.output/(name+ext)).write_text('stale')
        with self.assertLogs(level='WARNING'):
            self.assertEqual(inflation.generate(self.output,self.inf),[])
            self.assertEqual(consumer.generate(self.output,self.root/'missing'),[])
        self.assertEqual(list(self.output.iterdir()),[])
        with self.assertRaises(ValueError):rbi.discover_inflation(self.inf)

    def test_normal_india_generator_calls_both_canonical_modules(self):
        import generate_india as india
        from contextlib import ExitStack
        with ExitStack() as stack:
            stack.enter_context(patch('sys.argv',['generate_india.py']))
            stack.enter_context(patch.object(india,'OUTPUT_BASE',self.root/'generated'))
            stack.enter_context(patch.object(india,'load_india_data',return_value=None))
            stack.enter_context(patch.object(india,'load_forex_weekly',return_value=None))
            stack.enter_context(patch.object(india,'load_cag_data',return_value=(None,None,None)))
            for name in vars(india):
                if name.startswith('chart_'):stack.enter_context(patch.object(india,name))
            stack.enter_context(patch('charts.india_charts.india_cpi_contributions.generate'))
            inf=stack.enter_context(patch.object(inflation,'generate'))
            con=stack.enter_context(patch.object(consumer,'generate'))
            india.main()
            inf.assert_called_once();con.assert_called_once()
            self.assertEqual(inf.call_args,con.call_args)

    def test_canonical_insights_have_four_sections_and_no_release_values(self):
        for name in inflation.NAMES+consumer.NAMES:
            text=get_insight(name+'.png')
            sections = ['How to read this chart','Practical Takeaway','Frequency','Source']
            if name == consumer.NAMES[3]:
                sections = [section.upper() for section in sections]
            self.assertEqual([line.split(':**')[0][2:] for line in text.splitlines() if line.startswith('**')], sections)
            self.assertNotIn('2026',text)


if __name__=='__main__':unittest.main()

class DashboardPlacementTests(unittest.TestCase):
    def test_real_india_page_places_all_six_once_with_existing_sections(self):
        import ast
        from unittest.mock import MagicMock
        from charts.loader import chart_key
        root=Path(__file__).resolve().parents[1]
        tree=ast.parse((root/'app.py').read_text())
        page=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='page_india')
        code=compile(ast.Module(body=[page],type_ignores=[]),'app.py','exec')
        names=['01_india_pmi','03_india_fpi_monthly','06_india_inflation_bar',
               '14_india_credit_deposit','17_india_trade','07_india_expenditure_quality',
               *inflation.NAMES,*consumer.NAMES]
        charts=[Path(n+'.png') for n in names];sections=[];rendered=[]
        section_names=['Growth & Activity','Inflation','Monetary Conditions','External Sector',
                       'Public Finances','Equity Markets','Consumer Confidence']
        def header(index, answer):
            sections.append(section_names[index])
            return section_names[index]
        namespace={'st':MagicMock(),'get_charts':lambda _: (charts,'2026-10'),
                   'datetime':datetime,'_page_header_html':MagicMock(),
                   '_anchor':lambda _,title:title,'_load_soe':lambda _:None,
                   '_soe_changes_block':lambda _: '',
                   '_pop_summary':lambda paths,_:(None,paths),
                   'chart_key':chart_key,'_EH_ANCHORS':list(range(7)),'_EH_STYLE':'',
                   '_eh_india_answers':lambda *_:[None]*7,'_eh_india_observations':lambda:{},
                   '_eh_brief':lambda:{},'_eh_question_header':header,
                   '_render_iip_industry_heatmap':MagicMock(),'_render_cpi_items':MagicMock(),
                   '_render_india_equity_matrix':MagicMock(),
                   '_section':lambda title,**kwargs:sections.append(title),
                   '_render_grid':lambda paths,**kwargs:rendered.append((sections[-1],[p.stem for p in paths]))}
        exec(code,namespace);namespace['page_india']()
        self.assertEqual(sections,[section_names[i] for i in (0,1,2,5,3,4,6)])
        placed={}
        for section, items in rendered:
            placed.setdefault(section, []).extend(items)
        self.assertEqual(placed['Consumer Confidence'],
                         [consumer.NAMES[0], consumer.NAMES[3], *consumer.NAMES[1:3]])
        for name in inflation.NAMES:self.assertIn(name,placed['Inflation'])
        flat=[name for _,items in rendered for name in items]
        self.assertEqual(sorted(flat),sorted(names))
