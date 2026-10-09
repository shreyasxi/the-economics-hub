"""Median selection and approved canonical chart semantics."""
from datetime import datetime
import math
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
from data.processors import rbi_surveys as survey
from charts.india_charts import india_inflation_surveys as chart
from config.insights import get_insight


FIXTURE = Path(__file__).resolve().parent / 'fixtures/rbi_surveys/inflation.xlsx'

class HouseholdExpectationsTests(unittest.TestCase):
    def fixture(self, change=None):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = survey.SHEET
        for cell, value in {"B4": "Round No.", "C4": "Survey period ended",
                            "D5": "Current", "G5": "3 months ahead", "J5": "1 year ahead",
                            "E6": "Median ", "H6": "Median ", "K6": "Median "}.items():
            sheet[cell] = value
        for row, (round_no, date, values) in enumerate([
                ("79B", datetime(2025, 5, 1), (8.4, 9.5, 9.4)),
                (83, datetime(2026, 3, 1), (8.2, 9.3, 9.1)),
                ("83B", datetime(2026, 5, 1), (8.5, 9.7, 9.2))], 7):
            sheet.cell(row, 2, round_no)
            sheet.cell(row, 3, date)
            for column, value in zip((5, 8, 11), values):
                sheet.cell(row, column, value)
            # Distinct means prove that the parser selects only medians.
            for column in (4, 7, 10):
                sheet.cell(row, column, 99)
        # A separate standard-error row must never become an observation.
        sheet["E10"] = "[0.09]"
        sheet["H10"] = "[0.09]"
        sheet["K10"] = "[0.11]"
        if change:
            change(sheet)
        self.addCleanup(workbook.close)
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        path = Path(directory.name) / "fixture.xlsx"
        workbook.save(path)
        return path

    def test_median_selection_excludes_means_and_standard_errors(self):
        records, missing = survey.read_observations(self.fixture())
        self.assertEqual(len(records), 3)
        self.assertEqual(missing, [])
        latest, previous, earlier = survey.select_curves(records)
        self.assertEqual([r["round"] for r in (latest, previous, earlier)], ["83B", "83", "79B"])
        self.assertEqual([latest[k] for k in survey.SERIES], [8.5, 9.7, 9.2])
        self.assertEqual(set(latest), {"date", "round", "source_row", *survey.SERIES})

    def test_missing_values_stay_missing_and_comparisons_fail(self):
        records, missing = survey.read_observations(
            self.fixture(lambda sheet: setattr(sheet["H9"], "value", None)))
        self.assertEqual(missing, [{"row": 9, "series": "three_month"}])
        self.assertTrue(math.isnan(records[-1]["three_month"]))
        with self.assertRaisesRegex(ValueError, "Missing comparison data"):
            survey.select_curves(records)

    def test_duplicate_dates_bad_dates_non_numeric_and_changed_headers_rejected(self):
        for cell, value, message in [("C9", datetime(2026, 3, 1), "chronological"),
                                     ("C9", "May-26", "survey date"),
                                     ("K9", "[0.11]", "numeric"),
                                     ("E6", "Mean", "header")]:
            with self.subTest(cell=cell), self.assertRaisesRegex(ValueError, message):
                survey.read_observations(self.fixture(
                    lambda sheet: setattr(sheet[cell], "value", value)))

    @unittest.skipUnless(FIXTURE.exists(), "Official regression fixture is unavailable")
    def test_actual_workbook_all_three_comparison_medians(self):
        records, missing = survey.read_observations(FIXTURE)
        latest, previous, earlier = survey.select_curves(records)
        self.assertEqual(len(records), 91)
        # The workbook does not supply medians for its first four rounds.
        self.assertEqual(missing, [{"row": row, "series": name}
                                 for row in range(7, 11) for name in survey.SERIES])
        self.assertEqual([r["date"] for r in (latest, previous, earlier)],
                         [datetime(2026, 5, 1), datetime(2026, 3, 1), datetime(2025, 5, 1)])
        self.assertEqual(latest["source_row"], 143)
        self.assertEqual(latest["round"], "83B")
        for record, expected in zip((latest, previous, earlier), [
                (7.76373992885387, 9.32568168466109, 9.27841836756172),
                (7.20401865686928, 8.53106708397945, 8.79108430791818),
                (7.69755482717455, 8.88289076943664, 9.48054572560227)]):
            for name, value in zip(survey.SERIES, expected):
                self.assertAlmostEqual(record[name], value)

    def test_single_panel_two_line_header_and_latest_annotations(self):
        records, _ = survey.read_observations(self.fixture())
        with patch.object(chart.EconStyle, "save_chart") as save:
            chart.render_expectations(records, Path("unused.png"))
        fig = save.call_args.args[0]
        from matplotlib import pyplot as plt
        self.addCleanup(plt.close, fig)
        self.assertEqual(len(fig.axes), 1)
        ax = fig.axes[0]
        self.assertEqual(len(ax.lines), 3)
        self.assertEqual(ax.get_ylabel(), "Median inflation (%)")
        self.assertEqual([t.get_text() for t in fig.texts if t.get_position()[1] > .8],
                         [chart.TITLE, "Household perceived inflation and future expectations · Current vs. 3M vs. 1Y ahead"])
        self.assertEqual([t.get_text() for t in ax.texts], ["EXPECTATIONS", "8.50%", "9.70%", "9.20%"])
        self.assertEqual(chart.TITLE, "Inflation Expectations - India")
        self.assertEqual(tuple(fig.get_size_inches()), (9.2, 4.8))
        legend = fig.legends[0]
        self.assertFalse(legend.get_frame_on())
        self.assertTrue(all(t.get_fontsize() == 10 for t in legend.get_texts()))
        self.assertEqual([line.get_linestyle() for line in ax.lines], ["-", "--", "-"])
        self.assertEqual(len(ax.patches), 1)
        shade = ax.patches[0]
        vertices = ax.transData.inverted().transform(
            shade.get_transform().transform(shade.get_path().vertices))
        self.assertAlmostEqual(min(vertices[:, 0]), .5)
        self.assertAlmostEqual(max(vertices[:, 0]), 2.2)
        self.assertEqual(shade.get_alpha(), .8)

    def test_insights_exact_four_sections_and_official_source(self):
        records, _ = survey.read_observations(self.fixture())
        text = get_insight(chart.NAMES[0]+".png")
        self.assertEqual([line.split(":**")[0][2:] for line in text.splitlines() if line.startswith("**")],
                         ["How to read this chart", "Practical Takeaway", "Frequency", "Source"])
        self.assertIn("Bi-monthly", text)
        self.assertIn(survey.INFLATION_SOURCE, text)
        self.assertIn("BimonthlyPublications.aspx", survey.INFLATION_SOURCE)
        self.assertNotIn("slope", text.lower())
        self.assertNotIn(".xlsx", text)


if __name__ == "__main__":
    unittest.main()
